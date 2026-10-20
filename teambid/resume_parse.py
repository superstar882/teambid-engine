"""Resume file -> structured profile (facts only, nothing invented).

Pipeline: extract text (PDF via PyMuPDF, DOCX via python-docx, TXT/MD) -> split into sections by header vocabulary
-> parse experience entries around date ranges -> bullets (with wrapped-line repair) -> education, skills, projects,
certifications -> per-role stack and metrics -> years of experience, level and career lanes.

The parse is heuristic, so the UI shows it for review and every field is editable. With ANTHROPIC_API_KEY set,
llm.refine_profile() can clean it further (it may only restructure what is in the text).
"""
import datetime as dt
import re
from pathlib import Path

from . import skills as SK

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
MON = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\.?"
ONE_DATE = (rf"(?:\d{{1,2}}[./]\d{{1,2}}[./]\d{{4}}|\d{{4}}-\d{{2}}-\d{{2}}|(?:\d{{1,2}}\s+)?{MON}\s*,?\s*\d{{4}}|"
            rf"\d{{1,2}}[./]\d{{4}}|\d{{4}}-\d{{2}}|(?:Spring|Summer|Fall|Autumn|Winter)\s+\d{{4}}|\d{{4}})")
NOW = r"(?:Present|Current|Currently|Now|Today|Ongoing|To date|Till date|Until now|To now|Date)"
DATE_RANGE = re.compile(rf"(?i)({ONE_DATE})\s*(?:-|–|—|to|until|~)\s*({ONE_DATE}|{NOW})")
SINGLE_DATE = re.compile(rf"(?i)\b({MON}\s*,?\s*\d{{4}}|\d{{1,2}}/\d{{4}})\b")

SECTIONS = {
    "summary": ["summary", "professional summary", "profile", "about", "about me", "objective", "career objective", "overview", "career summary"],
    "experience": ["experience", "work experience", "professional experience", "employment", "employment history", "work history",
                   "relevant experience", "industry experience", "career history", "experience and projects"],
    "projects": ["projects", "personal projects", "selected projects", "side projects", "academic projects", "key projects", "portfolio"],
    "education": ["education", "academic background", "education and training", "academics"],
    "skills": ["skills", "technical skills", "core skills", "key skills", "skills and tools", "technologies", "tech stack",
               "tools", "competencies", "core competencies", "skills & tools", "technical proficiencies"],
    "certifications": ["certifications", "certificates", "licenses and certifications", "licenses & certifications", "courses",
                       "online courses", "training", "certifications and courses"],
    "awards": ["awards", "honors", "honors and awards", "achievements", "awards and honors", "accomplishments"],
    "languages": ["languages", "spoken languages"],
    "publications": ["publications", "papers", "research"],
    "volunteer": ["volunteer", "volunteering", "volunteer experience", "leadership", "activities", "leadership and activities", "extracurricular"],
    "interests": ["interests", "hobbies"],
}
HEADER_LOOKUP = {a: k for k, al in SECTIONS.items() for a in al}

TITLE_WORDS = re.compile(r"(?i)\b(engineer|developer|programmer|architect|manager|lead|intern|analyst|scientist|designer|consultant|"
                         r"specialist|administrator|director|founder|co-founder|cto|ceo|researcher|assistant|apprentice|contractor|"
                         r"associate|coordinator|officer|technician|head|vp|president|fellow|tutor|instructor|teacher|owner|"
                         r"freelancer|freelance|sre|devops|qa|tester|product|member of technical staff|mts|student)\b")
INTERN = re.compile(r"(?i)\b(intern|internship|co-op|coop|apprentice|trainee|student|research assistant|teaching assistant|"
                    r"undergraduate research|research apprentice|fellow)\b")
BULLET = re.compile(r"^\s*([•●▪◦‣∙·\-–—*►➤✓✔o])\s+")
LOCATION = re.compile(r"(?i)\b(remote|hybrid|[A-Z][a-zA-Z .'-]+,\s*(?:[A-Z]{2}|[A-Z][a-zA-Z ]+))\b")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"\+\d{1,3}[\s.-]?\(?\d{1,5}\)?(?:[\s.-]?\d{2,5}){2,4}|(?:\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
LINKEDIN = re.compile(r"(?i)(?:https?://)?(?:www\.)?linkedin\.com/in/[\w%-]+/?")
GITHUB = re.compile(r"(?i)(?:https?://)?(?:www\.)?github\.com/[\w-]+/?")
URL = re.compile(r"(?i)\b(?:https?://)?(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)*\.(?:com|io|dev|me|ai|app|net|org|co|xyz|site|tech)(?:/[\w./-]*)?\b")
METRIC = re.compile(r"(?i)(?<![A-Za-z])(?:[$€£]\s?\d[\d,.]*\s?(?:[kmb]|million|billion)?\+?|\d[\d,.]*\s?(?:%|x\b|\+|k\b|m\b|b\b|ms\b|s\b|"
                    r"hours?|hrs?|minutes?|mins?|days?|weeks?|months?|years?|users|customers|clients|requests|rows|records|"
                    r"files|lines|people|members|students|engineers|models|endpoints|services|tps|qps|rps)|\b\d{2,}[\d,.]*\b)")
DEGREE = re.compile(r"(?i)\b(Ph\.?\s?D\.?|Doctor(?:ate)?|M\.?B\.?A\.?|M\.?\s?S\.?c?\.?|M\.?\s?Eng\.?|M\.?\s?Tech|M\.?\s?A\.?|Master(?:'s)?(?: of [A-Z][a-z]+)?|"
                    r"B\.?\s?S\.?c?\.?|B\.?\s?A\.?|B\.?\s?Eng\.?|B\.?\s?Tech|B\.?\s?E\.?|Bachelor(?:'s)?(?: of [A-Z][a-z]+)?|"
                    r"Associate(?:'s)? (?:of|Degree)|A\.?A\.?S?\.?|High School Diploma|Diploma|Certificate)\b")
SCHOOL = re.compile(r"(?i)\b(university|college|institute|school|academy|polytechnic|universidad|universit[eéà])\b")
STATES = ("AL AK AZ AR CA CO CT DE FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR "
          "PA RI SC SD TN TX UT VT VA WA WV WI WY DC").split()
COUNTRIES = ("USA|United States|Canada|UK|United Kingdom|England|Scotland|Wales|Ireland|India|Japan|Singapore|Australia|"
             "New Zealand|Germany|France|Netherlands|Belgium|Luxembourg|Spain|Portugal|Italy|Malta|Greece|Cyprus|Austria|Switzerland|"
             "Denmark|Sweden|Norway|Finland|Iceland|Estonia|Latvia|Lithuania|Poland|Czechia|Czech Republic|Slovakia|Hungary|"
             "Romania|Bulgaria|Croatia|Slovenia|Serbia|Bosnia|Albania|North Macedonia|Montenegro|Ukraine|Moldova|Georgia|Armenia|"
             "Brazil|Mexico|Argentina|Colombia|Chile|Peru|Uruguay|Costa Rica|Israel|Turkey|UAE|United Arab Emirates|Saudi Arabia|"
             "Qatar|Egypt|Morocco|Tunisia|Nigeria|Kenya|Ghana|South Africa|China|Hong Kong|Taiwan|Korea|South Korea|Vietnam|"
             "Philippines|Myanmar|Thailand|Malaysia|Indonesia|Pakistan|Bangladesh|Sri Lanka|Nepal")
_CITY = r"[A-Z][a-zA-Z.'-]+(?: [A-Z][a-zA-Z.'-]+){0,2}"
PLACE = re.compile(r"\b(?:Remote|Hybrid|" + _CITY + r",\s*(?:" + "|".join(STATES) + r")\b|" + _CITY + r",\s*(?:" + COUNTRIES + r")\b)")
TITLE_START = re.compile(r"(?i)\b(senior|sr\.?|lead|staff|principal|junior|jr\.?|associate|head|chief|founding|founder|co-founder|"
                         r"software|data|machine|ml|ai|full[- ]?stack|front[- ]?end|back[- ]?end|research|undergraduate|graduate|"
                         r"engineering|product|devops|site|qa|member|applied|cloud|platform|mobile|ios|android|web|teaching|"
                         r"technical|solutions|security|systems|contract|freelance|intern|analytics|business|it|ux|ui|"
                         r"design|program|project|operations|marketing|sales|customer|support|network|database|embedded|game|"
                         r"working|student)\b")
GPA = re.compile(r"(?i)\bGPA[:\s]*([0-4]\.\d{1,2})(?:\s*/\s*([0-4]\.\d|\d{1,3}))?")


# ---------------------------------------------------------------- text extraction
def extract_text(path):
    """Return (text, meta) for a PDF, DOCX, TXT or MD resume. meta notes layout risks (columns, tables, images)."""
    p = Path(path)
    ext = p.suffix.lower()
    meta = {"pages": 1, "columns": False, "tables": False, "images": False, "file": p.name}
    if ext == ".pdf":
        import fitz
        doc = fitz.open(p)
        meta["pages"] = len(doc)
        parts = []
        for page in doc:
            blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
            w = page.rect.width
            right = [b for b in blocks if b[0] > w * 0.45 and len(b[4]) > 40]
            left = [b for b in blocks if b[2] < w * 0.55 and len(b[4]) > 40]
            if len(right) >= 3 and len(left) >= 3:
                meta["columns"] = True
            if page.get_images():
                meta["images"] = True
            parts.append(page.get_text("text", sort=True))
        return "\n".join(parts), meta
    if ext == ".docx":
        import docx
        d = docx.Document(p)
        lines = []
        for para in d.paragraphs:
            t = para.text
            if t.strip() and para.style is not None and "List" in (para.style.name or "") and not BULLET.match(t):
                t = "• " + t
            lines.append(t)
        if d.tables:
            meta["tables"] = True
            for tb in d.tables:
                for row in tb.rows:
                    cells = []
                    for c in row.cells:
                        if c.text.strip() and c.text not in cells:
                            cells.append(c.text)
                    lines.append(" | ".join(cells))
        for s in d.sections:
            for para in s.header.paragraphs:
                if para.text.strip():
                    lines.insert(0, para.text)
        return "\n".join(lines), meta
    return p.read_text(encoding="utf-8", errors="replace"), meta


# ---------------------------------------------------------------- helpers
def _norm_header(line):
    t = re.sub(r"[^a-z& ]", "", line.lower()).strip()
    t = re.sub(r"\s+", " ", t)
    return HEADER_LOOKUP.get(t) or HEADER_LOOKUP.get(t.replace("&", "and"))


def _is_header(line):
    s = line.strip().rstrip(":")
    if not s or len(s) > 45 or BULLET.match(line):
        return None
    key = _norm_header(s)
    if key:
        return key
    return None


def parse_date(s, end=False):
    """'Jan 2024' / '01/2024' / '2024' / 'Present' -> (year, month) or None."""
    if not s:
        return None
    s = s.strip().rstrip(".")
    if re.fullmatch(NOW, s, re.I):
        d = dt.date.today()
        return (d.year, d.month)
    m = re.match(r"(\d{1,2})([./])(\d{1,2})[./](\d{4})$", s)      # full date: day-first (EU) or month-first (US)
    if m:
        a, b, y = int(m[1]), int(m[3]), int(m[4])
        if a > 12:
            month = b                  # 15.03.2022 / 15/03/2022
        elif b > 12:
            month = a                  # 3/15/2022
        else:
            month = b if m[2] == "." else a   # dots are European (day.month.year); slashes default to US month/day
        return (y, max(1, min(12, month)))
    m = re.match(r"(\d{4})-(\d{2})-\d{2}$", s)
    if m:
        return (int(m[1]), int(m[2]))
    m = re.match(rf"(?i)(?:\d{{1,2}}\s+)?({MON})\s*,?\s*(\d{{4}})", s)
    if m:
        return (int(m[2]), MONTHS[m[1][:3].lower()])
    m = re.match(r"(\d{1,2})[./](\d{4})", s)                       # 03/2022, 03.2022
    if m:
        return (int(m[2]), max(1, min(12, int(m[1]))))
    m = re.match(r"(\d{4})-(\d{2})", s)
    if m:
        return (int(m[1]), int(m[2]))
    m = re.match(r"(?i)(spring|summer|fall|autumn|winter)\s+(\d{4})", s)
    if m:
        return (int(m[2]), {"spring": 3, "summer": 6, "fall": 9, "autumn": 9, "winter": 12}[m[1].lower()])
    m = re.match(r"(\d{4})", s)
    if m:
        return (int(m[1]), 12 if end else 1)
    return None


def months_between(a, b):
    if not a or not b:
        return 0
    return max(0, (b[0] - a[0]) * 12 + (b[1] - a[1]) + 1)


def _clean(line):
    line = line.replace(" ", " ").replace("\t", " ")
    return re.sub(r"\s{2,}", "  ", line).strip()


def metrics(text):
    out = []
    for m in METRIC.finditer(text):
        v = m.group(0).strip().rstrip(",.")
        if re.fullmatch(r"(19|20)\d\d", v):     # years are dates, not metrics
            continue
        if len(re.sub(r"\D", "", v)) == 0:
            continue
        out.append(v)
    return out


# ---------------------------------------------------------------- main
def parse(text, meta=None):
    meta = meta or {}
    raw_lines = [_clean(l) for l in text.splitlines()]
    lines = [l for l in raw_lines if l]
    notes = []

    # 1) split into sections
    sections, current, head = {}, "header", []
    for l in lines:
        key = _is_header(l)
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        if current == "header":
            head.append(l)
        else:
            sections.setdefault(current, []).append(l)
    if "experience" not in sections:
        notes.append("No 'Experience' header found: looked for jobs by dates across the whole resume.")
        sections["experience"] = [l for l in lines if l not in head[:3]]

    profile = {"name": "", "headline": "", "contact": {}, "summary": "", "roles": [], "projects": [], "education": [],
               "certifications": [], "awards": [], "languages": [], "skills": {}, "metrics": [], "parse_notes": notes,
               "source": meta}

    # 2) header block: name, contact, headline
    full = "\n".join(lines)
    c = profile["contact"]
    if (m := EMAIL.search(full)):
        c["email"] = m.group(0)
    if (m := LINKEDIN.search(full)):
        c["linkedin"] = m.group(0).rstrip("/")
    if (m := GITHUB.search(full)):
        c["github"] = m.group(0).rstrip("/")
    for l in head[:8] or lines[:8]:
        if "phone" not in c and (m := PHONE.search(l)):
            c["phone"] = m.group(0).strip()
        for u in URL.findall(EMAIL.sub(" ", l)):
            if "linkedin" not in u and "github" not in u and "website" not in c:
                c["website"] = u
        if "location" not in c and (loc := PLACE.search(l)) and loc.group(0) not in ("Remote", "Hybrid"):
            c["location"] = loc.group(0).strip()
    for l in (head or lines)[:4]:
        cand = re.sub(r"[|•·,].*$", "", l).strip()
        if 2 <= len(cand.split()) <= 4 and re.fullmatch(r"[A-Za-z][A-Za-z.'\- ]+", cand) and not TITLE_WORDS.search(cand):
            profile["name"] = cand.title() if cand.isupper() else cand
            break
    for l in head[:6]:
        if l.strip() == profile["name"] or EMAIL.search(l) or PHONE.search(l):
            continue
        if TITLE_WORDS.search(l) and len(l) < 120:
            profile["headline"] = l.strip(" |")
            break

    # 3) summary
    profile["summary"] = " ".join(sections.get("summary", []))[:1200]

    # 4) experience
    profile["roles"] = parse_entries(sections.get("experience", []))
    if not profile["roles"]:
        notes.append("Could not find jobs with dates. Check the Experience section uses date ranges like 'Jan 2023 - Present'.")

    # 5) projects (entries without required dates)
    profile["projects"] = parse_projects(sections.get("projects", []))

    # 6) education
    profile["education"] = parse_education(sections.get("education", []))

    # 7) skills: listed (skills section) vs proven (in bullets)
    listed_raw = []
    for l in sections.get("skills", []):
        body = l.split(":", 1)[1] if ":" in l and len(l.split(":", 1)[0]) < 40 else l
        for item in re.split(r"[,|;•·]|\s{2,}", body):
            item = item.strip(" .-–")
            if 1 < len(item) < 40:
                listed_raw.append(item)
    listed = SK.find("\n".join(sections.get("skills", [])) + "\n" + ", ".join(listed_raw), trusted=True)
    bullet_text = "\n".join(b["text"] for r in profile["roles"] for b in r["bullets"]) + "\n" + \
                  "\n".join(b["text"] for p in profile["projects"] for b in p["bullets"])
    meta_text = "\n".join(r.get("stack_line", "") for r in profile["roles"]) + "\n".join(p.get("stack_line", "") for p in profile["projects"])
    proven = SK.find(bullet_text + "\n" + meta_text, with_counts=True)
    all_found = SK.find(full, with_counts=True)
    labels = {v.lower() for v in SK.CATEGORY_LABEL.values()} | {
        "area", "technologies", "technology", "languages", "language", "frontend", "backend", "databases", "database", "tools",
        "cloud", "devops", "cloud & devops", "auth", "orms", "cms", "cms & e-commerce", "ui & animation", "testing", "skills",
        "frameworks", "libraries", "platforms", "other", "ui", "data", "ai", "ml", "mobile", "infrastructure", "methods"}
    unknown = [x for x in listed_raw if not SK.find(x, trusted=True) and x.lower() not in labels]
    profile["skills"] = {
        "proven": [s for s, _ in proven.most_common()],
        "listed_only": [s for s in listed if s not in proven],
        "other_listed": list(dict.fromkeys(unknown))[:40],
        "all": [s for s, _ in all_found.most_common()],
    }

    # 8) certifications, awards, languages
    profile["certifications"] = [l.lstrip("•-* ").strip() for l in sections.get("certifications", []) if len(l) > 3][:20]
    profile["awards"] = [l.lstrip("•-* ").strip() for l in sections.get("awards", []) if len(l) > 3][:20]
    langs = " ".join(sections.get("languages", []))
    profile["languages"] = [x.strip() for x in re.split(r"[,|;•]", langs) if x.strip()][:10]

    # 9) metrics (the only numbers tailoring may use), years, level, lanes
    profile["metrics"] = sorted(set(metrics(bullet_text + "\n" + profile["summary"])))
    derive(profile)
    return profile


def parse_entries(lines):
    """Experience lines -> roles. A role starts at a header line (non-bullet) that has a date range, together with
    the non-bullet lines right around it (company / title / location can be split over 1-3 lines)."""
    roles, pending, cur = [], [], None

    def is_bullet_like(l):
        return bool(BULLET.match(l)) or (len(l) > 90 and not DATE_RANGE.search(l) and not _is_stack_line(l))

    def flush_header(hdr_lines):
        nonlocal cur
        r = parse_role_header(hdr_lines)
        if r:
            roles.append(r)
            cur = r
            return True
        return False

    for l in lines:
        if _is_stack_line(l) and not BULLET.match(l):
            if pending and any(DATE_RANGE.search(p) for p in pending):
                flush_header(pending)
                pending = []
            if cur is not None:
                _attach_meta(cur, l)
            continue
        if is_bullet_like(l):
            if pending:
                if any(DATE_RANGE.search(p) for p in pending):
                    flush_header(pending)
                elif cur:
                    # stack lines like "Python | Spark | Kafka" sit between header and bullets
                    for p in pending:
                        _attach_meta(cur, p)
                pending = []
            if cur is None:
                continue
            text = BULLET.sub("", l).strip()
            cur["bullets"].append({"text": text})
            continue
        # non-bullet line
        if cur and cur["bullets"] and not DATE_RANGE.search(l) and _is_continuation(cur["bullets"][-1]["text"], l):
            cur["bullets"][-1]["text"] += " " + l.strip()
            continue
        pending.append(l)
        if DATE_RANGE.search(l):
            # include one more following header line (e.g. title under company) before flushing
            continue
        if len(pending) >= 2 and any(DATE_RANGE.search(p) for p in pending[:-1]):
            if not TITLE_WORDS.search(pending[-1]) and not LOCATION.fullmatch(pending[-1].strip()) and len(pending[-1]) > 60:
                flush_header(pending[:-1])
                pending = [pending[-1]]
            elif len(pending) >= 3:
                flush_header(pending)
                pending = []
    if pending and any(DATE_RANGE.search(p) for p in pending):
        flush_header(pending)
    for r in roles:
        finish_role(r)
    return roles


def _is_stack_line(line):
    return line.count("|") >= 2 and len(SK.find(line, trusted=True)) >= 2 and not DATE_RANGE.search(line)


def _split_company_title(part):
    """'Uber Software Engineer' -> ('Uber', 'Software Engineer'): the title starts at the first role word that is
    followed by a title noun (engineer, intern, analyst...)."""
    for m in TITLE_START.finditer(part):
        tail = part[m.start():]
        if m.start() > 0 and TITLE_WORDS.search(tail):
            return part[:m.start()].strip(" ,|-–—"), tail.strip(" ,|")
    return "", part


def _is_continuation(prev, line):
    if prev.endswith((".", "!", "?")) and line[:1].isupper():
        return False
    if TITLE_WORDS.search(line) and DATE_RANGE.search(line):
        return False
    return line[:1].islower() or not prev.endswith(".") and len(line) < 90 and not TITLE_WORDS.search(line[:40])


def _attach_meta(role, line):
    if len(SK.find(line, trusted=True)) >= 2 and len(line) < 220:
        role["stack_line"] = (role.get("stack_line", "") + " | " + line).strip(" |")


def parse_role_header(hdr):
    text = "  ".join(hdr)
    m = DATE_RANGE.search(text)
    if not m:
        return None
    start, end = parse_date(m[1]), parse_date(m[2], end=True)
    current = bool(re.fullmatch(NOW, m[2].strip(), re.I))
    rest = (text[:m.start()] + "  " + text[m.end():]).strip()
    location = ""
    lm = PLACE.search(rest)
    if lm:
        location = lm.group(0)
        rest = rest.replace(location, " ")
    parts = [p.strip(" ,|–—-·") for p in re.split(r"\s{2,}|\s[|–—·•]\s|\s-\s|\s@\s|\sat\s|\||·", rest)]
    parts = [p for p in parts if p and len(p) > 1 and not re.fullmatch(r"[\W\d]+", p)]
    title = next((p for p in parts if TITLE_WORDS.search(p)), "")
    company = next((p for p in parts if p != title), "")
    if title and not company:
        segs = [x.strip() for x in title.split(", ")]
        if len(segs) > 1 and TITLE_START.match(segs[0]) and TITLE_WORDS.search(segs[0]) and not TITLE_WORDS.search(segs[-1]):
            title, company = ", ".join(segs[:-1]), segs[-1]    # "Senior Data Engineer, Northwind" / "Working Student, Dev, Datenwerk"
        else:
            company, title = _split_company_title(title)      # "Uber Software Engineer"
    stack_line = ""
    if len(SK.find(company, trusted=True)) >= 2:           # "Python | Spark" mistaken for a company
        stack_line, company = company, next((p for p in parts if p not in (title, company)), "")
    return {"company": company, "title": title or (parts[1] if len(parts) > 1 else ""), "location": location.strip(),
            "start": f"{start[0]}-{start[1]:02d}" if start else "", "end": "Present" if current else (f"{end[0]}-{end[1]:02d}" if end else ""),
            "current": current, "months": months_between(start, end), "bullets": [], "stack_line": stack_line}


def finish_role(r):
    t = r["title"] + " " + r["company"]
    r["kind"] = "intern" if INTERN.search(t) else "job"
    for b in r["bullets"]:
        b["metrics"] = metrics(b["text"])
        b["skills"] = SK.find(b["text"])
    r["stack"] = list(dict.fromkeys(SK.find(r.get("stack_line", ""), trusted=True) + [s for b in r["bullets"] for s in b["skills"]]))


def parse_projects(lines):
    projects, cur = [], None
    for l in lines:
        if BULLET.match(l) or (cur and len(l) > 90):
            if cur is None:
                cur = {"name": "Project", "bullets": [], "stack_line": "", "link": ""}
                projects.append(cur)
            cur["bullets"].append({"text": BULLET.sub("", l).strip()})
            continue
        if cur and cur["bullets"] and _is_continuation(cur["bullets"][-1]["text"], l):
            cur["bullets"][-1]["text"] += " " + l
            continue
        name = re.split(r"\s[|–—-]\s|\s{2,}|:|\(", l)[0].strip()
        link = (URL.search(l) or [None])[0] if URL.search(l) else ""
        cur = {"name": name[:80], "bullets": [], "stack_line": l, "link": link or ""}
        projects.append(cur)
    for p in projects:
        for b in p["bullets"]:
            b["metrics"] = metrics(b["text"])
            b["skills"] = SK.find(b["text"])
        p["stack"] = list(dict.fromkeys(SK.find(p["stack_line"]) + [s for b in p["bullets"] for s in b["skills"]]))
    return [p for p in projects if p["bullets"] or len(p["stack"]) >= 1][:12]


def parse_education(lines):
    """Handles 'School | Degree, Field | 2016 - 2020 | GPA', 'School, B.Sc. in Field, 10.2016 - 09.2018' and
    '• Degree, Field · School' with the dates on the next line."""
    out, cur = [], None
    for raw in lines:
        l = BULLET.sub("", raw).strip()
        if not l:
            continue
        dr = DATE_RANGE.search(l)
        body = (l[:dr.start()] + " " + l[dr.end():]).strip(" ,|·") if dr else l
        segs = [x.strip(" ,|·-–—") for x in re.split(r"\s·\s|·|\s\|\s|\||\s{2,}|\s[–—]\s", body) if x.strip(" ,|·-–—")]
        school = next((x for x in segs if SCHOOL.search(x)), "")
        degree = next((x for x in segs if DEGREE.search(x)), "")
        if school and degree and school == degree:      # "School, B.Sc. in Field" in one piece: cut at the degree
            dm = DEGREE.search(school)
            school, degree = school[:dm.start()].strip(" ,"), school[dm.start():]
        if school or degree:
            if cur is None or (school and cur["school"]) or (degree and cur["degree"]):
                cur = {"school": "", "degree": "", "field": "", "start": "", "end": "", "gpa": "", "notes": []}
                out.append(cur)
            if school and not cur["school"]:
                cur["school"] = school
            if degree and not cur["degree"]:
                dm = DEGREE.search(degree)
                cur["degree"] = dm.group(0).strip()
                rest = degree[dm.end():]
                f = re.search(r"(?:\b(?:in|of)\s+|^[^,]*?,\s*|^\s+)([A-Z][A-Za-z &/]+?)\s*(?:,|\(|$|\d)", rest)
                if f and not re.fullmatch(r"(?i)degree|program|programme", f.group(1).strip()):
                    cur["field"] = f.group(1).strip()
        if cur is None:
            continue
        if (g := GPA.search(l)):
            cur["gpa"] = g.group(1) + (f"/{g.group(2)}" if g.group(2) else "")
        if dr and not cur["start"]:
            cur["start"], cur["end"] = dr[1], dr[2]
        elif not dr and not cur["end"] and (y := re.findall(r"\b(?:19|20)\d{2}\b", l)) and not school:
            cur["end"] = y[-1]
        if not school and not degree and not dr and not GPA.search(l) and len(l) > 3:
            cur["notes"].append(l)
    return [e for e in out if e["school"] or e["degree"]][:6]


def derive(profile):
    """Years of experience, level, career lanes and target titles from the parsed facts."""
    spans = []
    for r in profile["roles"]:
        if r["kind"] != "job" or not r["start"]:
            continue
        a = tuple(int(x) for x in r["start"].split("-"))
        b = dt.date.today() if r["current"] or not r["end"] else None
        b = (b.year, b.month) if b else tuple(int(x) for x in r["end"].split("-"))
        spans.append((a, b))
    months, last_end = 0, None
    for a, b in sorted(spans):                       # merge overlapping jobs
        if last_end and a <= last_end:
            if b > last_end:
                months += months_between(last_end, b) - 1
                last_end = b
            continue
        months += months_between(a, b)
        last_end = b
    intern_months = sum(r["months"] for r in profile["roles"] if r["kind"] == "intern")
    years = round(months / 12, 1)
    profile["years_experience"] = years
    profile["intern_months"] = intern_months
    titles = [r["title"] for r in profile["roles"]]
    senior_title = any(re.search(r"(?i)\b(senior|sr\.?|lead|staff|principal|manager|head|director)\b", t) for t in titles[:2])
    if years < 1:
        level = "entry"
    elif years < 3:
        level = "junior" if not senior_title else "mid"
    elif years < 6:
        level = "mid" if not senior_title or years < 4 else "senior"
    elif years < 10:
        level = "senior"
    else:
        level = "staff"
    profile["level"] = level
    from collections import Counter
    # recent work says most about where someone is heading: weight skills by how recent the role is
    counts = Counter({s: 1 for s in profile["skills"].get("listed_only", [])})
    for i, r in enumerate(profile["roles"]):
        w = 4 if i == 0 else 3 if i == 1 else 2 if i == 2 else 1
        for sk in r.get("stack", []):
            counts[sk] += w
    for pj in profile.get("projects", []):
        for sk in pj.get("stack", []):
            counts[sk] += 1
    lanes = SK.lane_scores(counts, titles[:3])
    ranked = [l for l, v in lanes.most_common() if v >= 3][:3] or ["backend"]
    profile["lanes"] = ranked
    tt = []
    for l in ranked:
        tt += SK.LANES[l][1][:3]
    profile["target_titles"] = list(dict.fromkeys(tt))[:8]
    return profile


def parse_file(path):
    text, meta = extract_text(path)
    prof = parse(text, meta)
    prof["raw_text"] = text
    return prof
